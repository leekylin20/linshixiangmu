from __future__ import annotations

import json
from pathlib import Path

from docx import Document


ROOT = Path(r"D:\BaiduNetdiskDownload\薛辉allin7月\md 薛辉")


def extract_text(path: Path) -> str:
    doc = Document(str(path))
    parts: list[str] = []
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def main() -> None:
    rows = []
    for path in sorted(ROOT.glob("*_原文.docx"), key=lambda p: p.name):
        text = extract_text(path)
        preview = text[:220].replace("\n", " ")
        rows.append(
            {
                "name": path.name,
                "size": path.stat().st_size,
                "chars": len(text),
                "paragraphs": text.count("\n") + (1 if text else 0),
                "preview": preview,
            }
        )
    print(json.dumps(rows, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
