from __future__ import annotations

import json
import re
import zipfile
from pathlib import Path

import docx
import openpyxl
import pypdf
import xlrd


ROOT = Path(r"D:\Program Files (x86)\微信记录\xwechat_files")
OUT = Path(r"E:\临时项目\wechat_extract")
OUT.mkdir(parents=True, exist_ok=True)

KEYWORDS = [
    "子木",
    "直播间",
    "直播",
    "报价",
    "价格",
    "设备",
    "灯光",
    "绿幕",
    "美颜",
    "机位",
    "采集卡",
    "云犀",
    "M4",
    "S3",
    "CPM",
    "CTR",
    "CVR",
    "ROI",
    "方案",
    "清单",
    "搭建",
    "参数",
]

EXTS = {".xlsx", ".xls", ".docx", ".pdf", ".txt", ".md", ".csv", ".json", ".html", ".htm", ".zip", ".rar"}


def relevant(path: Path) -> bool:
    text = path.name.lower()
    return path.suffix.lower() in EXTS and any(k.lower() in text for k in KEYWORDS)


def clean(s: object) -> str:
    if s is None:
        return ""
    text = str(s).replace("\r", " ").replace("\n", " ")
    text = re.sub(r"\s+", " ", text).strip()
    return repair_mojibake(text)


def repair_mojibake(text: str) -> str:
    if not text:
        return text
    markers = ("璁", "惧", "鐩", "鎶", "浠", "锛", "銆", "绾", "鏂", "寰")
    if not any(marker in text for marker in markers):
        return text
    try:
        repaired = text.encode("gb18030", errors="ignore").decode("utf-8", errors="ignore")
    except UnicodeError:
        return text
    if sum("\u4e00" <= ch <= "\u9fff" for ch in repaired) >= max(1, sum("\u4e00" <= ch <= "\u9fff" for ch in text) // 3):
        return repaired
    return text


def extract_xlsx(path: Path) -> str:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    chunks = []
    for ws in wb.worksheets:
        chunks.append(f"## Sheet: {ws.title}")
        shown = 0
        for row in ws.iter_rows(values_only=True):
            vals = [clean(v) for v in row]
            while vals and vals[-1] == "":
                vals.pop()
            if not any(vals):
                continue
            chunks.append("| " + " | ".join(vals) + " |")
            shown += 1
            if shown >= 120:
                chunks.append("...（该工作表仅截取前 120 行非空内容）")
                break
    return "\n".join(chunks)


def extract_xls(path: Path) -> str:
    book = xlrd.open_workbook(path)
    chunks = []
    for sheet in book.sheets():
        chunks.append(f"## Sheet: {sheet.name}")
        shown = 0
        for r in range(sheet.nrows):
            vals = [clean(sheet.cell_value(r, c)) for c in range(sheet.ncols)]
            while vals and vals[-1] == "":
                vals.pop()
            if not any(vals):
                continue
            chunks.append("| " + " | ".join(vals) + " |")
            shown += 1
            if shown >= 120:
                chunks.append("...（该工作表仅截取前 120 行非空内容）")
                break
    return "\n".join(chunks)


def extract_docx(path: Path) -> str:
    doc = docx.Document(str(path))
    chunks = []
    for p in doc.paragraphs:
        text = clean(p.text)
        if text:
            chunks.append(text)
    for i, table in enumerate(doc.tables, start=1):
        chunks.append(f"\n## Table {i}")
        for row in table.rows:
            vals = [clean(cell.text) for cell in row.cells]
            if any(vals):
                chunks.append("| " + " | ".join(vals) + " |")
    return "\n".join(chunks)


def extract_pdf(path: Path) -> str:
    reader = pypdf.PdfReader(str(path))
    chunks = []
    for i, page in enumerate(reader.pages[:30], start=1):
        text = clean(page.extract_text() or "")
        if text:
            chunks.append(f"## Page {i}\n{text}")
    if len(reader.pages) > 30:
        chunks.append(f"...（PDF 共 {len(reader.pages)} 页，仅截取前 30 页）")
    return "\n\n".join(chunks)


def extract_text(path: Path) -> str:
    for enc in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return path.read_text(encoding=enc, errors="strict")
        except UnicodeDecodeError:
            continue
    return path.read_text(encoding="utf-8", errors="ignore")


def extract_archive_listing(path: Path) -> str:
    if path.suffix.lower() == ".zip":
        with zipfile.ZipFile(path) as zf:
            return "\n".join(f"- {info.filename} ({info.file_size} bytes)" for info in zf.infolist())
    return "RAR 压缩包：当前脚本仅记录文件名，未解包读取内容。"


def extract(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        return extract_xlsx(path)
    if suffix == ".xls":
        return extract_xls(path)
    if suffix == ".docx":
        return extract_docx(path)
    if suffix == ".pdf":
        return extract_pdf(path)
    if suffix in {".txt", ".md", ".csv", ".json", ".html", ".htm"}:
        return extract_text(path)
    if suffix in {".zip", ".rar"}:
        return extract_archive_listing(path)
    return ""


def main() -> None:
    files = sorted(
        [p for p in ROOT.rglob("*") if p.is_file() and relevant(p)],
        key=lambda p: p.stat().st_mtime,
        reverse=True,
    )
    manifest = []
    combined = []
    for index, path in enumerate(files, start=1):
        rel = str(path.relative_to(ROOT))
        print(f"[{index}/{len(files)}] {rel}")
        item = {
            "path": str(path),
            "relative_path": rel,
            "size": path.stat().st_size,
            "mtime": path.stat().st_mtime,
            "status": "ok",
        }
        try:
            text = extract(path)
        except Exception as exc:
            item["status"] = "error"
            item["error"] = repr(exc)
            text = ""
        safe_name = re.sub(r'[<>:"/\\|?*\s]+', "_", rel)[:180]
        out_file = OUT / f"{index:03d}_{safe_name}.md"
        out_file.write_text(
            f"# {repair_mojibake(path.name)}\n\n来源：`{repair_mojibake(str(path))}`\n\n{text}\n",
            encoding="utf-8",
        )
        item["extracted_markdown"] = str(out_file)
        item["chars"] = len(text)
        manifest.append(item)
        if text:
            combined.append(f"# Source {index}: {path.name}\n\n来源：`{path}`\n\n{text}\n")
    (OUT / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    (OUT / "combined_relevant_sources.md").write_text("\n\n---\n\n".join(combined), encoding="utf-8")
    print(f"done: {len(files)} files, output {OUT}")


if __name__ == "__main__":
    main()
