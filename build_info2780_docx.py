from pathlib import Path
import os

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image


FOLDER = Path(os.environ["DOC_FOLDER"])
OUT = FOLDER / "【电商3C】GPT image2.0 详情页3.0全新工作流.docx"
SOURCE = "https://www.super-i.cn/info-2780.html"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text)
    r.bold = bold
    r.font.name = "Microsoft YaHei"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    r.font.size = Pt(10)


def add_para(doc, text="", style=None, bold=False, color=None):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    if text:
        r = p.add_run(text)
        r.bold = bold
        r.font.name = "Microsoft YaHei"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        if color:
            r.font.color.rgb = RGBColor(*color)
    return p


def add_bullets(doc, items):
    for item in items:
        p = add_para(doc, style="List Bullet")
        r = p.add_run(item)
        r.font.name = "Microsoft YaHei"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        r.font.size = Pt(10.5)


def add_numbered(doc, items):
    for item in items:
        p = add_para(doc, style="List Number")
        r = p.add_run(item)
        r.font.name = "Microsoft YaHei"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        r.font.size = Pt(10.5)


def add_code(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.4)
    p.paragraph_format.right_indent = Cm(0.2)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text)
    r.font.name = "Consolas"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    r.font.size = Pt(9.5)
    r.font.color.rgb = RGBColor(70, 70, 70)


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(6)
    for r in p.runs:
        r.font.name = "Microsoft YaHei"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        r.font.color.rgb = RGBColor(24, 74, 112) if level == 1 else RGBColor(46, 89, 132)
    return p


def add_image(doc, filename, caption, max_width_cm=14.8, max_height_cm=21.5):
    path = FOLDER / filename
    if not path.exists():
        return
    with Image.open(path) as im:
        w, h = im.size
    width_cm = min(max_width_cm, max_height_cm * (w / h))
    max_w = Cm(width_cm)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(path), width=max_w)
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(10)
    r = cap.add_run(caption)
    r.italic = True
    r.font.name = "Microsoft YaHei"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(90, 90, 90)


def style_document(doc):
    styles = doc.styles
    styles["Normal"].font.name = "Microsoft YaHei"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    styles["Normal"].font.size = Pt(10.5)
    for style_name, size in [("Heading 1", 18), ("Heading 2", 14), ("Heading 3", 12)]:
        style = styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True


def add_tool_table(doc):
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    headers = ["工具", "用途", "地址"]
    for i, h in enumerate(headers):
        set_cell_shading(table.rows[0].cells[i], "DDEBF7")
        set_cell_text(table.rows[0].cells[i], h, bold=True)
    rows = [
        ["Lovart 或 星流 AI", "视觉生成、画板素材管理、Agent 对话", "https://www.lovart.ai / https://www.xingliu.art"],
        ["豆包", "创建详情页策划智能体，拆解卖点与视觉脚本", "https://www.doubao.com"],
        ["GPT Image 2.0", "通过 Lovart 等工具调用，生成 Apple 风详情页画面", "工具内调用"],
    ]
    for row in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            set_cell_text(cells[i], text)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    doc.add_paragraph()


def build():
    doc = Document()
    style_document(doc)

    section = doc.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)

    title = "【电商3C】GPT image2.0 详情页3.0全新工作流"
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title)
    r.bold = True
    r.font.name = "Microsoft YaHei"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    r.font.size = Pt(22)
    r.font.color.rgb = RGBColor(24, 74, 112)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("一个人完成 Apple 风磁吸充电宝详情页策划、生成与微调")
    r.font.name = "Microsoft YaHei"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor(90, 90, 90)

    add_para(doc, f"来源：{SOURCE}")
    add_para(doc, "整理说明：本文档为页面教程的结构化整理版。网页中部分长提示词在源码中不可读，因此以可读取的步骤、工具、操作逻辑和图片素材为主进行排版。")

    add_image(doc, "01.jpg", "图 1：教程成品长图示例", max_width_cm=11.5)

    add_heading(doc, "教程概览", 1)
    add_para(doc, "本教程讲解一种“详情页生成 3.0”工作流：先用豆包大模型拆解产品卖点和详情页脚本，再把产品图、参数表和策划脚本输入 Lovart 或星流 AI，通过 GPT Image 2.0 生成 Apple 风 3C 产品详情页，并在最后进行局部微调。")
    add_bullets(doc, [
        "目标品类：3C 产品，示例为磁吸充电宝。",
        "目标风格：Apple 风、极简、高级、干净、科技感。",
        "核心流程：素材准备 → 豆包策划智能体 → Lovart 视觉生成 → 高阶微调 → 输出长图。",
    ])

    add_heading(doc, "准备工具", 1)
    add_tool_table(doc)

    add_heading(doc, "第一步：素材与对标准备", 1)
    add_para(doc, "正式生成前，需要先准备两类输入材料：一份高质量对标详情页，以及自己的产品资料。对标图用于让 AI 理解版式、视觉节奏和风格标准；产品资料用于确保生成内容不偏离真实卖点。")
    add_numbered(doc, [
        "对标详情页：选择你认为排版优秀、风格高级的竞品详情页长图。网页端展示不便时，可把长图裁切后分段使用。",
        "产品信息：准备产品主图、场景图、关键参数表和卖点信息。",
    ])
    add_image(doc, "02.jpg", "图 2：对标详情页长图，用于参考排版与视觉节奏", max_width_cm=10.5)
    add_image(doc, "03.png", "图 3：产品图素材", max_width_cm=12.8)
    add_image(doc, "04.jpg", "图 4：产品参数表素材", max_width_cm=14.5)

    add_heading(doc, "第二步：用豆包创建专属策划智能体", 1)
    add_para(doc, "打开豆包，创建一个面向详情页策划的智能体。它的任务不是直接画图，而是先把产品信息拆成详情页模块、卖点结构、视觉文案和每屏脚本。")
    add_bullets(doc, [
        "先上传对标详情页，让智能体分析页面风格、模块顺序、文案层级和视觉语言。",
        "再上传自己的产品图和参数表，让智能体提炼产品卖点。",
        "最后要求智能体输出完整的 10 大模块文案与视觉脚本。",
    ])
    add_para(doc, "页面推荐采用分步“投喂”方式，因为该模式一次只能上传一张图：")
    add_code(doc, "输入 1：上传对标详情页图片，请分析它的版式结构、视觉风格、模块顺序和文案层级。\n输入 2：上传产品图片，请识别产品外观特征和可用于详情页呈现的视觉重点。\n输入 3：上传参数图片，请提炼核心规格、功能卖点和消费者关心的信息。\n输入 4：请基于以上材料，生成一套完整的 10 屏详情页模块文案与视觉脚本。")
    add_image(doc, "05.jpg", "图 5：豆包创建智能体与操作入口", max_width_cm=14.5)
    add_image(doc, "06.png", "图 6：豆包输出的策划智能体提示词示例", max_width_cm=14.5)
    add_image(doc, "07.png", "图 7：上传对标详情页后的分析示例", max_width_cm=14.5)
    add_image(doc, "08.png", "图 8：上传产品参数后的信息提炼示例", max_width_cm=14.5)
    add_image(doc, "09.png", "图 9：详情页 10 大模块文案与视觉脚本输出", max_width_cm=14.5)

    add_heading(doc, "第三步：Lovart 视觉生成", 1)
    add_para(doc, "拿到豆包生成的模块脚本后，进入 Lovart 或星流 AI，把对标图、产品图、参数图和整理好的脚本拖入画板，再与 Agent 对话生成详情页画面。")
    add_bullets(doc, [
        "先生成 5 张，用于测试 AI 是否理解 Apple 风的质感、字体、留白和排版。",
        "如果方向正确，再继续要求它按同一标准生成后 5 屏。",
        "生成时要强调白底、极简、科技感、产品质感、模块化长图和电商详情页语境。",
    ])
    add_code(doc, "核心生成思路：请基于上传的产品图、参数表和详情页脚本，生成 Apple 风 3C 产品详情页。要求白底极简、高级科技感、清晰信息层级、产品主体突出、文案排版精致，并按 10 屏详情页模块输出。")
    add_image(doc, "10.png", "图 10：在 Lovart 画板中导入素材并开始生成", max_width_cm=14.8)
    add_image(doc, "11.png", "图 11：将豆包生成的脚本输入 Lovart Agent", max_width_cm=14.8)
    add_image(doc, "12.png", "图 12：首轮生成的详情页模块预览", max_width_cm=14.8)
    add_image(doc, "13.png", "图 13：前 5 屏详情页生成结果", max_width_cm=14.8)

    add_heading(doc, "第四步：高阶微调", 1)
    add_para(doc, "AI 直出的图通常能达到可用水平，但还需要人工导演式微调。教程重点提到两个常见问题：画面太素，以及首屏冲击力不足。")
    add_heading(doc, "问题 1：画面太素，缺乏层次", 2)
    add_para(doc, "极简风容易被 AI 做成大面积死白，看起来像未完成草图。可以选择单调的图片重新生成，并要求在不破坏白底和整体高级感的前提下增加色彩层次。")
    add_code(doc, "微调指令：给这张图增加多一些色彩，现在有些寡淡；白底不要变，颜色不要太乱。")
    add_image(doc, "14.png", "图 14：针对画面寡淡问题的微调指令", max_width_cm=14.8)
    add_image(doc, "15.png", "图 15：增加层次后的详情页局部效果", max_width_cm=14.8)

    add_heading(doc, "问题 2：首屏缺乏视觉冲击力", 2)
    add_para(doc, "详情页第一屏决定用户是否继续浏览。若首屏过于平铺直叙，应单独优化产品角度、空间感、近景冲击力和高端品牌感。")
    add_code(doc, "微调思路：单独优化首屏，让产品主体更有视觉冲击力；增强近景透视、光影质感和品牌级留白，保持 Apple 风极简高级。")
    add_image(doc, "16.png", "图 16：首屏冲击力优化结果", max_width_cm=10.5)

    add_heading(doc, "最终输出", 1)
    add_para(doc, "完成微调后，将各屏拼接或导出为完整详情页长图。最终结果应具备清晰卖点结构、统一视觉风格、干净留白和明确的电商转化信息。")
    add_image(doc, "17.png", "图 17：最终长图输出示例 1", max_width_cm=4.2)
    add_image(doc, "18.png", "图 18：最终长图输出示例 2", max_width_cm=4.2)

    add_heading(doc, "总结", 1)
    add_para(doc, "这套“详情页生成 3.0”工作流的关键，不是直接让 AI 画整张图，而是先让豆包完成卖点拆解和脚本策划，再把脚本交给 Lovart / GPT Image 2.0 做视觉落地。最后由创作者针对色彩、层次、首屏冲击力和信息表达进行微调。")
    add_bullets(doc, [
        "豆包负责策划：拆卖点、排模块、写文案、定视觉脚本。",
        "Lovart / GPT Image 2.0 负责生成：把脚本转化成视觉详情页。",
        "人工负责导演：筛选、补色、增强首屏、统一风格。",
    ])

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
