from pathlib import Path
import html
import os
import re
import shutil
import tempfile
import zipfile

from bs4 import BeautifulSoup
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt
from mobi import extract as mobi_extract
from pypdf import PdfReader


SRC = Path(os.environ["BOOK_SRC"])
OUT = Path(os.environ["DOCX_OUT"])
OUT.mkdir(parents=True, exist_ok=True)


def safe_name(name: str, limit: int = 110) -> str:
    name = re.sub(r'[<>:"/\\|?*\r\n]+', "_", name).strip(" .")
    return name[:limit].rstrip(" .")


def normalize_text(text: str) -> str:
    text = html.unescape(text or "")
    text = "".join(
        ch for ch in text
        if ch in "\n\r\t" or ord(ch) >= 32
    )
    text = text.replace("\u3000", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def iter_epub_html(path: Path):
    with zipfile.ZipFile(path) as zf:
        names = [n for n in zf.namelist() if n.lower().endswith((".html", ".xhtml", ".htm"))]
        names.sort()
        for name in names:
            try:
                data = zf.read(name)
            except Exception:
                continue
            yield name, data


def html_to_blocks(data: bytes):
    soup = BeautifulSoup(data, "html.parser")
    for tag in soup(["script", "style", "noscript"]):
        tag.decompose()
    blocks = []
    for node in soup.find_all(["h1", "h2", "h3", "h4", "p", "li", "blockquote"]):
        text = normalize_text(node.get_text(" ", strip=True))
        if not text:
            continue
        level = None
        if node.name in {"h1", "h2", "h3", "h4"}:
            level = min(int(node.name[1]), 4)
        elif node.name == "li":
            text = "- " + text
        blocks.append((level, text))
    if not blocks:
        text = normalize_text(soup.get_text("\n", strip=True))
        blocks = [(None, p.strip()) for p in text.splitlines() if p.strip()]
    return blocks


def extract_epub(path: Path):
    for _, data in iter_epub_html(path):
        for block in html_to_blocks(data):
            yield block


def find_html_files(root: Path):
    return sorted([p for p in root.rglob("*") if p.is_file() and p.suffix.lower() in {".html", ".xhtml", ".htm"}])


def extract_mobi_like(path: Path):
    tmp = Path(tempfile.mkdtemp(prefix="book_unpack_"))
    try:
        unpack_dir, _ = mobi_extract(str(path))
        unpack_dir = Path(unpack_dir)
        html_files = find_html_files(unpack_dir)
        for html_file in html_files:
            try:
                data = html_file.read_bytes()
            except Exception:
                continue
            for block in html_to_blocks(data):
                yield block
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def extract_pdf_text(path: Path):
    reader = PdfReader(str(path))
    for idx, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception:
            text = ""
        text = normalize_text(text)
        if text:
            yield (2, f"第 {idx} 页")
            for para in re.split(r"\n\s*\n|(?<=。)\s+", text):
                para = normalize_text(para)
                if para:
                    yield (None, para)


def setup_doc(title: str):
    doc = Document()
    styles = doc.styles
    styles["Normal"].font.name = "Microsoft YaHei"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    styles["Normal"].font.size = Pt(10.5)
    for name in ["Title", "Heading 1", "Heading 2", "Heading 3", "Heading 4"]:
        style = styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    doc.add_heading(title, 0)
    return doc


def add_blocks_to_doc(doc: Document, blocks):
    count = 0
    for level, text in blocks:
        if not text:
            continue
        if len(text) > 3000:
            parts = re.findall(r".{1,1800}(?:。|$)", text)
        else:
            parts = [text]
        for part in parts:
            part = part.strip()
            if not part:
                continue
            if level:
                doc.add_heading(part[:240], level=min(level, 4))
            else:
                doc.add_paragraph(part)
            count += 1
    return count


def convert_one(path: Path):
    out_path = OUT / (safe_name(path.stem) + ".docx")
    title = path.stem
    doc = setup_doc(title)
    suffix = path.suffix.lower()
    if suffix == ".epub":
        blocks = extract_epub(path)
    elif suffix in {".mobi", ".azw3"}:
        blocks = extract_mobi_like(path)
    elif suffix == ".pdf":
        blocks = extract_pdf_text(path)
    else:
        return None
    count = add_blocks_to_doc(doc, blocks)
    doc.save(out_path)
    return out_path, count


def main():
    targets = []
    for p in SRC.iterdir():
        if not p.is_file():
            continue
        if p.suffix.lower() != ".pdf":
            targets.append(p)
        elif "神经科学原理" in p.name:
            targets.append(p)
    for p in sorted(targets, key=lambda x: (x.suffix.lower(), x.name)):
        try:
            result = convert_one(p)
            if result:
                out_path, count = result
                print(f"OK\t{p.name}\t{count}\t{out_path}")
        except Exception as exc:
            print(f"FAIL\t{p.name}\t{type(exc).__name__}: {exc}")


if __name__ == "__main__":
    main()
