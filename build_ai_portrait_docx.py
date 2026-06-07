from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image


OUT = Path(r"E:\临时项目\AI 人像提示词的本质.docx")
ASSET_DIR = Path(r"E:\临时项目\AI人像提示词的本质_素材")


def set_font(run, size=10.5, bold=False, color=None):
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_paragraph(doc, text=""):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.18
    p.paragraph_format.space_after = Pt(7)
    if text:
        r = p.add_run(text)
        set_font(r)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    for r in p.runs:
        set_font(r, size=15 if level == 1 else 12, bold=True, color=(31, 78, 121))
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(item)
        set_font(r)


def add_prompt_box(doc, lines):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    cell.text = ""
    for i, line in enumerate(lines):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(line)
        set_font(r, size=9.5, color=(60, 60, 60))
    doc.add_paragraph()


def add_image(doc, filename, caption):
    path = ASSET_DIR / filename
    with Image.open(path) as im:
        w, h = im.size
    width_cm = min(14.2, 12.7 * (w / h))
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(str(path), width=Cm(width_cm))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(10)
    r = cap.add_run(caption)
    set_font(r, size=9, color=(95, 95, 95))
    r.italic = True


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)

    style = doc.styles["Normal"]
    style.font.name = "Microsoft YaHei"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    style.font.size = Pt(10.5)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("AI 人像提示词的本质")
    set_font(r, size=22, bold=True, color=(31, 78, 121))

    add_paragraph(doc, "很多人在生成 AI 人像时，会把重点放在“漂亮”“真实”“电影感”这类结果词上。但这些词本身太抽象，模型会倾向于生成训练集中最常见、最安全、最模板化的脸。")
    add_paragraph(doc, "真正决定人像是否自然的，不是单独写一句“realistic portrait”，而是把人像拆成摄影能看见的具体信息：皮肤纹理、面部微瑕疵、年龄痕迹、光线方向、镜头焦段、景深、妆容状态和环境色温。")

    add_heading(doc, "问题：只写结果词，容易得到塑料感", 1)
    add_paragraph(doc, "如果提示词只描述“一个美丽女性的人像照片”，模型通常会生成高度平均化的脸：皮肤过度光滑、五官比例完美但缺少个人特征，整体看起来像广告模板或 AI 默认美颜。")
    add_prompt_box(doc, [
        "基础思路：",
        "a portrait of a beautiful young woman, realistic photo"
    ])
    add_image(doc, "01.png", "基础提示词生成效果：五官整齐，但皮肤和面部细节偏模板化")

    add_heading(doc, "改进：把“真实”拆成可见细节", 1)
    add_paragraph(doc, "更有效的写法，是让模型看到一个具体的人，而不是一个抽象的“美女”。人像提示词应该补充年龄、肤质、光线、镜头和不完美细节。")
    add_bullets(doc, [
        "年龄与身份：不要只写 young woman，可以写 28 岁、亚洲女性、自然神态。",
        "皮肤质感：加入 pores、freckles、subtle blemishes、natural skin texture 等细节方向。",
        "光线：写清楚 window light、soft side lighting、natural daylight 等真实摄影条件。",
        "镜头：加入 85mm lens、shallow depth of field、portrait photography 等摄影语言。",
        "不完美：适度加入轻微黑眼圈、细小斑点、真实肤色变化，让画面脱离美颜模板。",
    ])
    add_prompt_box(doc, [
        "增强思路：",
        "28 岁女性人像，自然表情，真实皮肤纹理，轻微雀斑与毛孔，柔和窗边侧光，浅景深，85mm 人像镜头，摄影棚外的自然日光质感。"
    ])
    add_image(doc, "02.png", "增强提示词生成效果：皮肤、光线和面部特征更接近真实摄影")

    add_heading(doc, "核心结论", 1)
    add_paragraph(doc, "AI 人像提示词的本质，不是堆叠“真实、漂亮、高清、电影感”，而是把抽象审美翻译成模型能执行的视觉条件。")
    add_paragraph(doc, "当提示词描述的是“摄影机实际能拍到什么”，而不是“我希望它看起来怎样”，人像才会更自然、更有个体差异，也更接近真实照片。")

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
